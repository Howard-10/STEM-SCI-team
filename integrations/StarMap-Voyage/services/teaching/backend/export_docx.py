"""Markdown -> DOCX 导出工具（基于 python-docx）。

支持：标题(#/##/###)、加粗、行内代码、无序/有序列表、表格、代码块、普通段落。
用于把生成的教学设计 markdown 转成可下载的 Word 文档。
"""
import re
import io
from typing import List, Optional

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH

CN_FONT = "微软雅黑"
MONO_FONT = "Consolas"


def _set_run_font(run, size: int = 11, bold: bool = False, color: Optional[tuple] = None, mono: bool = False):
    name = MONO_FONT if mono else CN_FONT
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def _add_inline(paragraph, text: str, base_size: int = 11):
    """解析 **bold** 与 `code` 行内标记，追加 run。"""
    tokens = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) >= 4:
            run = paragraph.add_run(tok[2:-2])
            _set_run_font(run, base_size, bold=True)
        elif tok.startswith("`") and tok.endswith("`") and len(tok) >= 2:
            run = paragraph.add_run(tok[1:-1])
            _set_run_font(run, base_size - 1, mono=True, color=(79, 70, 229))
        else:
            run = paragraph.add_run(tok)
            _set_run_font(run, base_size)


def _add_heading(doc: Document, text: str, level: int):
    p = doc.add_paragraph()
    p.space_before = Pt(14 if level == 1 else 10)
    size = 20 if level == 1 else (16 if level == 2 else 13)
    color = (49, 46, 129) if level <= 2 else (67, 56, 202)
    run = p.add_run(text)
    _set_run_font(run, size, bold=True, color=color)
    return p


def _add_code_block(doc: Document, code: str):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), "F5F5F5")
    pPr.append(shd)
    for line in code.split("\n"):
        if line != code.split("\n")[0]:
            p.add_run().add_break()
        run = p.add_run(line)
        _set_run_font(run, 9, mono=True)


def _flush_table(doc: Document, rows: List[List[str]]):
    if not rows:
        return
    # 去掉 markdown 表头分隔行（--- 或 :--: 之类）
    data = []
    for r in rows:
        if all(re.match(r"^:?-{2,}:?$", c.strip()) for c in r if c.strip()):
            continue
        data.append(r)
    if not data:
        return
    ncols = max(len(r) for r in data)
    table = doc.add_table(rows=len(data), cols=ncols)
    try:
        table.style = "Table Grid"
    except Exception:
        pass
    for ri, row in enumerate(data):
        for ci in range(ncols):
            cell = table.cell(ri, ci)
            text = row[ci].strip() if ci < len(row) else ""
            cell.text = ""
            para = cell.paragraphs[0]
            _add_inline(para, text, base_size=10)
            for run in para.runs:
                _set_run_font(run, 10)


def markdown_to_docx(markdown: str, title: Optional[str] = None) -> Document:
    doc = Document()
    lines = markdown.split("\n")

    # 文档主标题
    if title:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(title)
        _set_run_font(run, 22, bold=True, color=(49, 46, 129))
        doc.add_paragraph()

    i = 0
    table_rows: List[List[str]] = []
    code_buf: List[str] = []
    in_code = False

    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        # 代码块
        if stripped.startswith("```"):
            if in_code:
                _add_code_block(doc, "\n".join(code_buf))
                code_buf = []
                in_code = False
            else:
                in_code = True
                code_buf = []
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        # 表格
        if stripped.startswith("|"):
            cells = [c.strip() for c in stripped.strip("|").split("|")]
            table_rows.append(cells)
            i += 1
            continue
        else:
            _flush_table(doc, table_rows)
            table_rows = []

        # 标题
        m = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if m:
            _add_heading(doc, m.group(2).strip(), len(m.group(1)))
            i += 1
            continue

        # 分隔线 / 空行
        if stripped == "" or re.match(r"^-{3,}$", stripped):
            i += 1
            continue

        # 无序列表
        m = re.match(r"^[-*]\s+(.+)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            _add_inline(p, m.group(1))
            i += 1
            continue

        # 有序列表
        m = re.match(r"^\d+[.、)]\s*(.+)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Number")
            _add_inline(p, m.group(1))
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        _add_inline(p, stripped)
        i += 1

    _flush_table(doc, table_rows)
    if in_code and code_buf:
        _add_code_block(doc, "\n".join(code_buf))

    return doc


def markdown_to_docx_bytes(markdown: str, title: Optional[str] = None) -> bytes:
    doc = markdown_to_docx(markdown, title)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
